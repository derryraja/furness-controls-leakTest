"""
Excel results log: a styled ``Results`` sheet (one row per test) and a
``Summary`` sheet with live formulas and charts, rebuilt on every save.
"""

from __future__ import annotations

import os

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .runner import TestResult

C_HEADER_BG, C_HEADER_FG = "1F4E79", "FFFFFF"
C_PASS_BG, C_PASS_FG = "C6EFCE", "276221"
C_FAIL_BG, C_FAIL_FG = "FFC7CE", "9C0006"
C_ALT_ROW, C_TITLE_FG = "DDEEFF", "1F4E79"
C_STAT_BG, C_NUM_BG = "EBF3FB", "D9E8F5"

THIN = Side(style="thin", color="AAAAAA")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLUMNS = [
    ("#", 6),
    ("Test Reference", 18),
    ("Device Number", 16),
    ("Date / Time", 22),
    ("Operator", 14),
    ("Start (mbar)", 14),
    ("End (mbar)", 12),
    ("Drop (mbar)", 12),
    ("Leak Rate\n(mbar/s)", 16),
    ("Threshold\n(mbar/s)", 16),
    ("Result", 10),
]
COL_NUM, COL_LEAK, COL_RESULT = 1, 9, 11
N_COLS = len(COLUMNS)
LR_COL = get_column_letter(COL_LEAK)
RES_COL = get_column_letter(COL_RESULT)
FIRST_DATA_ROW = 3

TITLE = "Leak Test Log"
SUMMARY_TITLE = "Leak Test Summary"


def _build_summary_sheet(wb, threshold: float):
    if "Summary" in wb.sheetnames:
        del wb["Summary"]
    ws = wb.create_sheet("Summary", 0)
    ws.sheet_view.showGridLines = False

    for col, w in {"A": 2, "B": 28, "C": 20, "D": 4, "E": 18, "F": 18, "G": 4}.items():
        ws.column_dimensions[col].width = w
    for r in range(1, 45):
        ws.row_dimensions[r].height = 22

    ws.merge_cells("B1:F1")
    t = ws["B1"]
    t.value = SUMMARY_TITLE
    t.font = Font(bold=True, name="Arial", size=14, color="FFFFFF")
    t.fill = PatternFill("solid", fgColor=C_HEADER_BG)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 34

    rng = lambda col: f"Results!{col}{FIRST_DATA_ROW}:{col}100000"
    stats = [
        ("Total Tests", f"=COUNTA({rng('B')})", None),
        ("Passes", f'=COUNTIF({rng(RES_COL)},"PASS")', None),
        ("Failures", f'=COUNTIF({rng(RES_COL)},"FAIL")', None),
        ("Pass %", "=IFERROR(C3/C2,0)", "0.0%"),
        ("Fail %", "=IFERROR(C4/C2,0)", "0.0%"),
        ("Max Leak Rate (mbar/s)", f"=IFERROR(MAX({rng(LR_COL)}),0)", "0.00000"),
        ("Min Leak Rate (mbar/s)", f"=IFERROR(MIN({rng(LR_COL)}),0)", "0.00000"),
        ("Avg Leak Rate (mbar/s)", f"=IFERROR(AVERAGE({rng(LR_COL)}),0)", "0.00000"),
    ]
    for i, (label, formula, fmt) in enumerate(stats, start=2):
        lc = ws.cell(row=i, column=2, value=label)
        lc.font = Font(bold=True, name="Arial", size=11, color=C_TITLE_FG)
        lc.fill = PatternFill("solid", fgColor=C_NUM_BG)
        lc.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        lc.border = BORDER
        vc = ws.cell(row=i, column=3, value=formula)
        vc.font = Font(bold=True, name="Arial", size=11)
        vc.fill = PatternFill("solid", fgColor=C_STAT_BG)
        vc.alignment = Alignment(horizontal="center", vertical="center")
        vc.border = BORDER
        if fmt:
            vc.number_format = fmt

    ws["C3"].fill = PatternFill("solid", fgColor=C_PASS_BG)
    ws["C3"].font = Font(bold=True, name="Arial", size=11, color=C_PASS_FG)
    ws["C4"].fill = PatternFill("solid", fgColor=C_FAIL_BG)
    ws["C4"].font = Font(bold=True, name="Arial", size=11, color=C_FAIL_FG)
    ws["B10"].value = f"Current pass/fail threshold: {threshold} mbar/s"
    ws["B10"].font = Font(italic=True, name="Arial", size=9, color="666666")

    # Helper cells feeding the charts
    for cell, val in {"E2": "Pass", "F2": "=C3", "E3": "Fail", "F3": "=C4",
                      "E5": "Max", "F5": "=C7", "E6": "Min", "F6": "=C8",
                      "E7": "Avg", "F7": "=C9"}.items():
        ws[cell] = val
        ws[cell].font = Font(name="Arial", size=9, color="888888")

    def colour_points(series, colours):
        pts = []
        for idx, (fill, line) in enumerate(colours):
            pt = DataPoint(idx=idx)
            pt.graphicalProperties.solidFill = fill
            if line:
                pt.graphicalProperties.line.solidFill = line
            pts.append(pt)
        series.dPt = pts

    bar = BarChart()
    bar.type, bar.title, bar.style = "col", "Pass vs Fail", 2
    bar.y_axis.title = "Number of tests"
    bar.width, bar.height = 14, 10
    bar.add_data(Reference(ws, min_col=6, min_row=2, max_row=3))
    bar.set_categories(Reference(ws, min_col=5, min_row=2, max_row=3))
    bar.legend = None
    colour_points(bar.series[0], [("4CAF50", "276221"), ("F44336", "9C0006")])
    ws.add_chart(bar, "B12")

    pie = PieChart()
    pie.title, pie.style, pie.width, pie.height = "Pass % vs Fail %", 10, 14, 10
    pie.add_data(Reference(ws, min_col=6, min_row=2, max_row=3))
    pie.set_categories(Reference(ws, min_col=5, min_row=2, max_row=3))
    pie.series[0].dLbls = DataLabelList()
    pie.series[0].dLbls.showPercent = True
    pie.series[0].dLbls.showCatName = True
    pie.series[0].dLbls.showVal = False
    colour_points(pie.series[0], [("4CAF50", None), ("F44336", None)])
    ws.add_chart(pie, "E12")

    lr = BarChart()
    lr.type, lr.title, lr.style = "col", "Leak Rate Stats (mbar/s)", 2
    lr.y_axis.title = "mbar/s"
    lr.width, lr.height = 14, 10
    lr.add_data(Reference(ws, min_col=6, min_row=5, max_row=7))
    lr.set_categories(Reference(ws, min_col=5, min_row=5, max_row=7))
    lr.legend = None
    colour_points(lr.series[0], [("2196F3", "0D47A1"), ("FF9800", "E65100"), ("9C27B0", "4A148C")])
    ws.add_chart(lr, "B28")


def _new_workbook(threshold: float) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.freeze_panes = f"A{FIRST_DATA_ROW}"
    ws.sheet_view.showGridLines = False

    ws.merge_cells(f"A1:{get_column_letter(N_COLS)}1")
    tc = ws["A1"]
    tc.value = TITLE
    tc.font = Font(bold=True, color=C_TITLE_FG, name="Arial", size=13)
    tc.alignment = Alignment(horizontal="center", vertical="center")
    tc.fill = PatternFill("solid", fgColor=C_NUM_BG)
    ws.row_dimensions[1].height = 28

    for col_idx, (label, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=2, column=col_idx, value=label)
        cell.font = Font(bold=True, color=C_HEADER_FG, name="Arial", size=10)
        cell.fill = PatternFill("solid", fgColor=C_HEADER_BG)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[2].height = 36

    _build_summary_sheet(wb, threshold)
    return wb


def _style_row(ws, row: int, result: str):
    bg = C_ALT_ROW if row % 2 == 0 else "FFFFFF"
    for col in range(1, N_COLS + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = PatternFill("solid", fgColor=bg)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
        cell.font = Font(name="Arial", size=10)
    nc = ws.cell(row=row, column=COL_NUM)
    nc.font = Font(bold=True, name="Arial", size=10, color=C_TITLE_FG)
    nc.fill = PatternFill("solid", fgColor=C_NUM_BG)
    rc = ws.cell(row=row, column=COL_RESULT)
    ok = result == "PASS"
    rc.fill = PatternFill("solid", fgColor=C_PASS_BG if ok else C_FAIL_BG)
    rc.font = Font(bold=True, color=C_PASS_FG if ok else C_FAIL_FG, name="Arial", size=10)


def _last_data_row(ws) -> int:
    """Last row with a test reference (ignores formatted-but-empty rows)."""
    row = ws.max_row
    while row >= FIRST_DATA_ROW and ws.cell(row=row, column=2).value in (None, ""):
        row -= 1
    return row


def log_result(path: str, r: TestResult) -> int:
    """Append one result. Returns the entry number. Raises PermissionError
    if the workbook is open in Excel."""
    if os.path.exists(path):
        wb = load_workbook(path)
        ws = wb["Results"]
    else:
        folder = os.path.dirname(os.path.abspath(path))
        os.makedirs(folder, exist_ok=True)
        wb = _new_workbook(r.threshold)
        ws = wb["Results"]

    next_row = max(_last_data_row(ws) + 1, FIRST_DATA_ROW)
    entry = next_row - FIRST_DATA_ROW + 1
    values = [
        entry, r.test_reference, r.device_number,
        r.timestamp.strftime("%Y-%m-%d  %H:%M:%S"), r.operator,
        r.pressure_start, r.pressure_end, r.pressure_drop,
        r.leak_rate, r.threshold, r.result,
    ]
    for col_idx, val in enumerate(values, start=1):
        ws.cell(row=next_row, column=col_idx, value=val)
    _style_row(ws, next_row, r.result)
    _build_summary_sheet(wb, r.threshold)
    wb.save(path)
    return entry


def read_recent(path: str, limit: int = 50) -> list[tuple]:
    """Return the last ``limit`` rows of the Results sheet (newest first)."""
    if not os.path.exists(path):
        return []
    wb = load_workbook(path, read_only=True)
    try:
        ws = wb["Results"]
        rows = [r for r in ws.iter_rows(min_row=FIRST_DATA_ROW, values_only=True)
                if r and r[1] not in (None, "")]
    finally:
        wb.close()
    return list(reversed(rows[-limit:]))
